#!/usr/bin/env python3
"""把 markdown 路线文档转成排版良好的 Word (.docx)。"""
import re
import sys
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

SRC = sys.argv[1] if len(sys.argv) > 1 else "/home/user/ai_engineer_plan/AI工程师学习路线.md"
OUT = sys.argv[2] if len(sys.argv) > 2 else "/home/user/ai_engineer_plan/AI工程师学习路线.docx"
FONT = "微软雅黑"

doc = Document()

# 全局字体（含中文）
style = doc.styles["Normal"]
style.font.name = FONT
style.font.size = Pt(10.5)
style.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
style.paragraph_format.space_after = Pt(4)
style.paragraph_format.line_spacing = 1.25

for i, size in enumerate([22, 16, 13, 11.5], start=0):
    try:
        h = doc.styles[f"Heading {i}" if i else "Title"]
    except KeyError:
        continue
    h.font.name = FONT
    h.font.size = Pt(size)
    h.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
    try:
        h.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    except AttributeError:
        pass

# 页边距
for sec in doc.sections:
    sec.top_margin = sec.bottom_margin = Cm(2.2)
    sec.left_margin = sec.right_margin = Cm(2.2)


def add_runs(par, text):
    """处理 **加粗** 和 `代码`。"""
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = par.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = par.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0xB0, 0x30, 0x30)
        else:
            par.add_run(part)


def clean(text):
    text = text.replace("★", "★ ")
    return text.strip()


lines = open(SRC, encoding="utf-8").read().split("\n")
i = 0
n_tables = n_code = 0
while i < len(lines):
    line = lines[i]

    # 代码块
    if line.strip().startswith("```"):
        i += 1
        buf = []
        while i < len(lines) and not lines[i].strip().startswith("```"):
            buf.append(lines[i])
            i += 1
        i += 1
        if buf:
            n_code += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.space_before = Pt(4)
            r = p.add_run("\n".join(buf))
            r.font.name = "Consolas"
            r.font.size = Pt(9)
            r.font.color.rgb = RGBColor(0x22, 0x44, 0x22)
        continue

    # 表格
    if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:\-\|]+\|$", lines[i + 1]):
        block = []
        while i < len(lines) and lines[i].startswith("|"):
            if not re.match(r"^\|[\s:\-\|]+\|$", lines[i]):
                block.append([c.strip() for c in lines[i].strip().strip("|").split("|")])
            i += 1
        if block:
            n_tables += 1
            ncols = max(len(r) for r in block)
            t = doc.add_table(rows=0, cols=ncols)
            t.style = "Light Grid Accent 1"
            for ri, row in enumerate(block):
                cells = t.add_row().cells
                for ci in range(ncols):
                    txt = row[ci] if ci < len(row) else ""
                    par = cells[ci].paragraphs[0]
                    par.paragraph_format.space_after = Pt(0)
                    add_runs(par, txt)
                    for r in par.runs:
                        r.font.size = Pt(9)
                        if ri == 0:
                            r.bold = True
                            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                    if ri == 0:
                        sh = cells[ci]._tc.get_or_add_tcPr()
                        el = sh.makeelement(qn("w:shd"), {qn("w:fill"): "2E5C8A"})
                        sh.append(el)
            doc.add_paragraph()
        continue

    s = line.strip()
    if not s:
        i += 1
        continue
    if s.startswith("---"):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(2)
        r = p.add_run("─" * 46)
        r.font.color.rgb = RGBColor(0xBB, 0xBB, 0xBB)
        i += 1
        continue

    m = re.match(r"^(#{1,4})\s+(.*)$", s)
    if m:
        level, text = len(m.group(1)), clean(m.group(2))
        if level == 1:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(text)
            r.bold = True
            r.font.size = Pt(20)
            r.font.color.rgb = RGBColor(0x1F, 0x38, 0x64)
            r.font.name = FONT
        else:
            doc.add_heading(text, level=min(level - 1, 3))
        i += 1
        continue

    if s.startswith(">"):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.6)
        add_runs(p, clean(s.lstrip("> ")))
        for r in p.runs:
            r.italic = True
            r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        i += 1
        continue

    m = re.match(r"^[-*]\s+(.*)$", s)
    if m:
        p = doc.add_paragraph(style="List Bullet")
        add_runs(p, clean(m.group(1)))
        i += 1
        continue

    m = re.match(r"^(\d+)\.\s+(.*)$", s)
    if m:
        p = doc.add_paragraph(style="List Number")
        add_runs(p, clean(m.group(2)))
        i += 1
        continue

    p = doc.add_paragraph()
    add_runs(p, clean(s))
    i += 1

doc.save(OUT)
print(f"✅ 已生成: {OUT}")
print(f"   段落 {len(doc.paragraphs)} ｜ 表格 {n_tables} ｜ 代码块 {n_code}")
